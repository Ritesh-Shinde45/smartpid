from io import BytesIO

import pandas as pd
from openpyxl import Workbook, load_workbook

from core.allocator import allocate
from core.exporter import export_workbook
from core.models import load_config
from core.tq_generator import build_tqs
from core.validator import validate, violations_to_df


def _export(pipeline, **kw):
    res = pipeline["result"]
    viol = violations_to_df(validate(res.table, pipeline["rules"], pipeline["hardware"]))
    tqs = build_tqs(res.unplaced, pipeline["unmapped"])
    return export_workbook(res.table, res.stats, viol, tqs, pipeline["hardware"], project="Demo", **kw), tqs


def test_workbook_has_all_sheets_and_rows(pipeline):
    data, _ = _export(pipeline)
    wb = load_workbook(BytesIO(data))
    assert wb.sheetnames == ["Summary", "Allocation", "Cabinet Map", "Violations", "TQ List"]
    assert wb["Allocation"].max_row == len(pipeline["result"].table) + 1


def test_cabinet_map_contains_every_tag(pipeline):
    wb = load_workbook(BytesIO(_export(pipeline)[0]))
    cells = {c.value for row in wb["Cabinet Map"].iter_rows() for c in row if c.value}
    assert set(pipeline["result"].table["tag"]) <= cells


def test_tq_for_unmapped_types(pipeline):
    _, tqs = _export(pipeline)
    assert set(tqs["category"]) == {"Unknown signal type"} and len(tqs) == 2
    assert tqs["tq_id"].tolist() == ["TQ-001", "TQ-002"]


def test_tq_for_oversized_cable_and_llm_fallback():
    df = pd.DataFrame([{"tag": f"T{i}", "description": "", "signal_type": "AI", "charm_type": "AI",
                        "cable_id": "BIG", "ab_role": "", "pair_id": ""} for i in range(90)])
    res = allocate(df, load_config("default_rules.json"), load_config("hardware.json"))
    assert build_tqs(res.unplaced)["category"].iloc[0] == "Cable exceeds CIOC capacity"

    def boom(rows):
        raise RuntimeError("api down")
    assert len(build_tqs(res.unplaced, llm_fn=boom)) == 1      # LLM failure is harmless


def test_template_sheets_are_kept(pipeline, tmp_path):
    tpl = tmp_path / "tpl.xlsx"
    wb = Workbook(); wb.active.title = "Cover"; wb.active["A1"] = "Customer cover page"; wb.save(tpl)
    wb2 = load_workbook(BytesIO(_export(pipeline, template_path=str(tpl))[0]))
    assert wb2["Cover"]["A1"].value == "Customer cover page" and "Allocation" in wb2.sheetnames