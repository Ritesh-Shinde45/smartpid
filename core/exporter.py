from __future__ import annotations

from datetime import datetime
from io import BytesIO
from typing import Optional, Union

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from core.models import Hardware

HEADER_FILL = PatternFill("solid", fgColor="1F3864")
HEADER_FONT = Font(bold=True, color="FFFFFF")
SPARE_FILL = PatternFill("solid", fgColor="F2F2F2")
CHARM_COLORS = {"AI": "BDD7EE", "AO": "F8CBAD", "DI": "C6E0B4", "DO": "FFE699", "RTD": "D9D2E9", "TC": "F4B183"}
SHEETS = ["Summary", "Allocation", "Cabinet Map", "Violations", "TQ List"]


def _header(ws, row: int, values: list) -> None:
    for c, v in enumerate(values, 1):
        cell = ws.cell(row=row, column=c, value=v)
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _autofit(ws, cap: int = 45) -> None:
    for col in ws.columns:
        width = max((len(str(c.value)) for c in col if c.value is not None), default=8)
        ws.column_dimensions[get_column_letter(col[0].column)].width = min(max(width + 2, 8), cap)


def _write_df(ws, df: pd.DataFrame) -> None:
    _header(ws, 1, list(df.columns))
    for r, rec in enumerate(df.itertuples(index=False), 2):
        for c, v in enumerate(rec, 1):
            ws.cell(row=r, column=c, value=None if pd.isna(v) else v)
    ws.freeze_panes = "A2"
    if len(df.columns):
        ws.auto_filter.ref = f"A1:{get_column_letter(len(df.columns))}{max(len(df) + 1, 2)}"
    _autofit(ws)


def _summary(ws, table, stats, hw, violations, tqs, project, rev) -> None:
    ws["A1"], ws["A1"].font = "SmartPID - IO Allocation Summary", Font(bold=True, size=14)
    info = [
        ("Project", project), ("Revision", rev), ("Generated", datetime.now().strftime("%Y-%m-%d %H:%M")),
        ("Signals in", stats.get("signals_in", len(table))), ("Signals allocated", len(table)),
        ("Signals not allocated (see TQ List)", stats.get("signals_unplaced", 0)),
        ("Cabinets used", stats.get("cabinets_used", 0)),
        ("Rule violations", len(violations)), ("Open TQs", len(tqs)),
    ]
    for i, (k, v) in enumerate(info, 3):
        ws.cell(row=i, column=1, value=k).font = Font(bold=True)
        ws.cell(row=i, column=2, value=v)

    row = len(info) + 5
    _header(ws, row, ["Cabinet", "Used CHARMs", "Capacity", "Fill %", "Free CHARMs"])
    for cab, d in sorted(stats.get("fill_by_cabinet", {}).items()):
        row += 1
        for c, v in enumerate([f"MC-{cab:02d}", d["used"], d["capacity"], d["fill_pct"], d["capacity"] - d["used"]], 1):
            ws.cell(row=row, column=c, value=v)

    if not table.empty:
        row += 2
        _header(ws, row, ["CHARM type", "Signals"])
        for ctype, n in table["charm_type"].value_counts().items():
            row += 1
            ws.cell(row=row, column=1, value=ctype)
            ws.cell(row=row, column=2, value=int(n))
    _autofit(ws)


def _cabinet_map(ws, table: pd.DataFrame, hw: Hardware) -> None:
    lookup = {(r.cabinet, r.cioc, r.baseplate, r.charm): (r.tag, r.charm_type) for r in table.itertuples()}
    row = 1
    for cab in sorted(table["cabinet"].unique()) if not table.empty else []:
        ws.cell(row=row, column=1, value=f"MC-{int(cab):02d}").font = Font(bold=True, size=13)
        row += 1
        _header(ws, row, ["CIOC", "Baseplate"] + [f"CH{i}" for i in range(1, hw.charms_per_baseplate + 1)])
        for ci in range(1, hw.cioc_per_cabinet + 1):
            for bp in range(1, hw.baseplates_per_cioc + 1):
                row += 1
                ws.cell(row=row, column=1, value=f"CIOC-{ci}")
                ws.cell(row=row, column=2, value=f"BP-{bp}")
                for ch in range(1, hw.charms_per_baseplate + 1):
                    cell = ws.cell(row=row, column=2 + ch)
                    hit = lookup.get((cab, ci, bp, ch))
                    if hit:
                        cell.value = hit[0]
                        cell.fill = PatternFill("solid", fgColor=CHARM_COLORS.get(hit[1], "E7E6E6"))
                    else:
                        cell.fill = SPARE_FILL
        row += 2
    ws.freeze_panes = "C1"
    _autofit(ws, cap=18)


def export_workbook(
    table: pd.DataFrame,
    stats: dict,
    violations_df: pd.DataFrame,
    tqs: pd.DataFrame,
    hardware: Union[dict, Hardware],
    project: str = "",
    rev: str = "A",
    template_path: Optional[str] = None,
) -> bytes:
    hw = hardware if isinstance(hardware, Hardware) else Hardware.from_dict(hardware)
    wb = load_workbook(template_path) if template_path else Workbook()
    if not template_path:
        wb.remove(wb.active)
    for name in SHEETS:
        if name in wb.sheetnames:
            del wb[name]

    _summary(wb.create_sheet("Summary"), table, stats, hw, violations_df, tqs, project, rev)
    _write_df(wb.create_sheet("Allocation"), table)
    _cabinet_map(wb.create_sheet("Cabinet Map"), table, hw)
    _write_df(wb.create_sheet("Violations"), violations_df)
    _write_df(wb.create_sheet("TQ List"), tqs)

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()